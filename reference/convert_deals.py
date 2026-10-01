# -*- coding: utf-8 -*-
"""
convert_deals.py – Chuyển file export Deals từ CRM (raw) sang file template phân tích.

Cách chạy:
    python convert_deals.py <file_raw.xlsx> <file_output.xlsx> [--ref-date dd/mm/yyyy]

--ref-date: ngày chốt số liệu, dùng để kiểm tra "Close date đã qua".
            Mặc định lấy từ tên file (..._dd_mm_yy.xlsx), nếu không có thì lấy ngày hôm nay.

Yêu cầu: pandas, openpyxl.
"""
import sys, re, html, unicodedata, argparse
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

# ============================== THAM SỐ ==============================
USD_MISLABEL_THRESHOLD = 1_000_000   # deal gắn USD có giá trị >= ngưỡng này coi là VND nhập nhầm nhãn
USD_RATE = 26_000                    # GIẢ ĐỊNH, chỉ dùng cho deal USD < ngưỡng (hiện chưa có deal nào)
BOM_TOLERANCE = 0.05                 # lệch giá trị so với tổng BOM > 5% thì ghi chú
DEAL_SIZE = [(1e9, '< 1B'), (2e9, '1B - 2B'), (5e9, '2B - 5B')]   # >= 5B -> '> 5B'
TEST_NAMES = re.compile(r'(^|[^a-z])test([^a-z]|$)')

# Cột bắt buộc phải có trong file raw (kiểm tra trước khi chạy)
REQUIRED = ['ID', 'Tên thương vụ', 'Giá trị thương vụ', 'Tiền tệ thương vụ', 'Chủ sở hữu thương vụ',
            'Giai đoạn', 'Trạng thái thương vụ', 'Tên công ty', 'Bảng danh mục sản phẩm dự án',
            'Phân loại HĐ', 'Reseller', 'Tên hãng', 'Thông tin BOM (model, số lượng, năm support...)',
            'Tên dự án', 'Timeline dự án', 'Ngày ký hợp đồng', 'Phân loại nhóm deal',
            'Lý do thất bại', 'Failed additional comment', 'Từ lúc', 'Cập nhật lần cuối']

VENDOR_MAP = {
    'KASPERSKY': 'Kaspersky', 'KAPERSKY': 'Kaspersky', 'GTB': 'GTB', 'SOPHOS': 'Sophos', 'OPSWAT': 'OPSWAT',
    'NETSCOUT': 'NetScout', 'DELINEA': 'Delinea', 'CLOUDFLARE': 'Cloudflare', 'PROGRESS': 'Progress',
    'INFOEXPRESS': 'Infoexpress', 'OPENTEXT': 'OpenText', 'BARRACUDA': 'Barracuda', 'RADWARE': 'Radware',
    'SAFETICA': 'Safetica', 'SAFEBREACH': 'SafeBreach', 'HCL SOFTWARE': 'HCL Software', 'HCL': 'HCL Software',
    'ZECURION': 'Zecurion', 'ACRONIS': 'Acronis', 'QUALYS': 'Qualys', 'STELLAR CYBER': 'Stellar Cyber',
    'PENTA SECURITY': 'Penta Security', 'ARCON': 'Arcon', 'PENTERA': 'Pentera', 'PENTARA': 'Pentera',
    'TRUSTWAVE': 'Trustwave', 'NETGEAR': 'Netgear', 'THALES': 'Thales', 'FORCEPOINT': 'Forcepoint',
    'SECPOD': 'SecPod', 'PROOFPOINT': 'Proofpoint', 'CYBLE': 'Cyble', 'FORTRA': 'Fortra', 'VIAVI': 'Viavi',
    'TRELLIX': 'Trellix', 'HÃNG KHÁC': 'Khác', 'SẢN PHẨM KHÁC': 'Khác', 'KHÁC': 'Khác',
}
# Từ khóa để suy ra hãng khi BOM không có "Hãng:" (tìm trên text đã bỏ dấu, chữ thường)
VENDOR_KEYWORDS = [
    ('Kaspersky', r'kaspersky|kapersky|\bkas\b'), ('Sophos', r'sophos'), ('GTB', r'\bgtb\b'),
    ('OPSWAT', r'opswat|metadefender'), ('NetScout', r'netscout|arbor'), ('Delinea', r'delinea'),
    ('Cloudflare', r'cloudflare'), ('Progress', r'progress|whatsup|\bwug\b|loadmaster|kemp'),
    ('Infoexpress', r'infoexpress|easynac'), ('OpenText', r'opentext|arcsight|fortify|\bsmax\b|\bopds\b|voltage|encase'),
    ('Barracuda', r'barracuda'), ('Radware', r'radware'), ('Safetica', r'safetica'), ('SafeBreach', r'safebreach'),
    ('HCL Software', r'\bhcl\b|appscan|\bhlc\b'), ('Zecurion', r'zecurion'), ('Acronis', r'acronis'),
    ('Qualys', r'qualys'), ('Stellar Cyber', r'stellar'), ('Penta Security', r'penta'), ('Arcon', r'arcon'),
    ('Pentera', r'pentera|pentara'), ('Trellix', r'trellix'), ('SecPod', r'secpod'), ('Thales', r'thales'),
    ('Proofpoint', r'proofpoint'), ('Forcepoint', r'forcepoint'), ('Trustwave', r'trustwave'),
    ('Netgear', r'netgear'), ('Viavi', r'viavi'),
]
# Từ khóa nhận diện phần "hãng/giải pháp" trong tên thương vụ
SOLUTION_KW = {'waf', 'ddos', 'dr', 'fw', 'firewall', 'sophos', 'kaspersky', 'kapersky', 'kas', 'gtb', 'safetica',
               'safebreach', 'smax', 'loadmaster', 'opswat', 'qualys', 'progress', 'pam', 'secpod', 'renew',
               'foundations', 'optimum', 'appscan', 'hlc', 'endpoint', 'dlp', 'nac', 'easynac', 'trellix', 'delinea',
               'netscout', 'opds', 'cloudflare', 'opentext', 'barracuda', 'radware', 'infoexpress', 'acronis',
               'zecurion', 'stellar', 'pentera', 'fortify', 'arcsight', 'wug', 'whatsup', 'hcl', 'kemp', 'edr',
               'siem', 'mfa', 'ztna', 'appsec', 'db', 'security'}
FILLER = {'va', 'and', 'plus', 'cap', 'do', 'bo', 'sung'}
PROJECT_PREFIX = re.compile(r'^(mua sam|giai phap|du an|trien khai|xay dung|nang cap|thay the)')


# ============================== HÀM TIỆN ÍCH ==============================
def fold(s):
    """chữ thường, bỏ dấu tiếng Việt – dùng để so khớp"""
    s = unicodedata.normalize('NFD', str(s).lower()).replace('đ', 'd')
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def clean_text(x):
    if not isinstance(x, str):
        return x
    x = html.unescape(x)
    x = re.sub(r'<[^>]+>', ' ', x)          # bỏ thẻ HTML (<p>…)
    x = re.sub(r'\s+', ' ', x).strip()
    return x or None


def parse_date(s):
    if isinstance(s, (pd.Timestamp, datetime)):
        return pd.Timestamp(s)
    if not isinstance(s, str) or not s.strip():
        return pd.NaT
    s = s.strip().replace('-', '/')
    for f in ('%d/%m/%Y %H:%M', '%d/%m/%Y'):
        try:
            return pd.to_datetime(s, format=f)
        except ValueError:
            pass
    return pd.NaT


def vnd(x):
    return f"{x:,.0f}".replace(',', '.')


def to_num(s):
    try:
        return float(s.replace(',', ''))
    except (ValueError, AttributeError):
        return None


BOM_SPLIT = re.compile(r'\s*\|?\s*,\s*(?=(?:Hãng:|Số lượng:|Tên giải pháp:))')


def parse_bom(s):
    """Tách cột 'Bảng danh mục sản phẩm dự án' thành các dòng {vendor, amount}"""
    if not isinstance(s, str):
        return []
    items = []
    for part in BOM_SPLIT.split(s):
        part = part.strip()
        if not part:
            continue
        m = re.search(r'Hãng:\s*([^|]*?)\s*(?:\||$)', part)
        vendor = (m.group(1).strip() or None) if m else None
        t = re.search(r'Thành tiền:\s*([\d,\.]+)', part)
        amt = to_num(t.group(1)) if t else None
        if amt is None:   # thiếu Thành tiền nhưng có Số lượng x Đơn giá
            q = re.search(r'Số lượng:\s*([\d,\.]+)', part)
            p = re.search(r'Đơn giá:\s*([\d,\.]+)', part)
            if q and p and to_num(q.group(1)) is not None and to_num(p.group(1)) is not None:
                amt = to_num(q.group(1)) * to_num(p.group(1))
        items.append({'vendor': vendor, 'amount': amt})
    return items


def canon_vendor(v):
    return VENDOR_MAP.get(v.strip().upper(), v.strip().title()) if v else None


def is_solution(seg):
    toks = [t for t in re.split(r'[^a-z0-9]+', fold(seg)) if t]
    if not toks:
        return False
    hit = sum(1 for t in toks if t in SOLUTION_KW or t in FILLER or t.isdigit())
    return hit >= max(1, len(toks) * 0.6)


def customer_from_deal_name(name):
    """Trả về (tên khách hàng, cách lấy)"""
    n = name.strip().rstrip(';').strip()
    f = fold(n)
    if TEST_NAMES.search(f):
        return n, 'test'
    if re.match(r'^C12-BCA 175 - ', n):                         # phần sau là reseller
        return 'C12-BCA 175', 'tach'
    mc = re.search(r'\bcho\s+(.+)$', n)
    if PROJECT_PREFIX.match(f) and mc:                          # "Mua sắm … cho X" -> X
        return mc.group(1).strip(), 'tach'
    if '_' in n:                                                # "KhachHang_Hang/GiaiPhap"
        parts = [p.strip() for p in n.split('_') if p.strip()]
        keep = [p for p in parts if not is_solution(p)]
        if keep and len(keep) < len(parts):
            return ' - '.join(keep), 'tach'
        if len(keep) == len(parts):
            return ' - '.join(parts), 'giu'
        return n, 'kiem tra'
    m = re.match(r'^(.*\S)\s*-\s*([^-]+)$', n)                  # "KhachHang - Hang"
    if m and is_solution(m.group(2)):
        return m.group(1).strip(), 'tach'
    if PROJECT_PREFIX.match(f):
        return n, 'kiem tra'
    return n, 'giu'


def deal_size(v):
    if v is None or pd.isna(v) or v <= 0:
        return None
    for lim, lab in DEAL_SIZE:
        if v < lim:
            return lab
    return '> 5B'


def vendor_info(row):
    """Trả về (hãng chính, [các hãng], cách xác định, tổng BOM, số dòng BOM chưa có giá)"""
    items = parse_bom(row['Bảng danh mục sản phẩm dự án'])
    last, amt, cnt, order = None, {}, {}, []
    for it in items:
        v = canon_vendor(it['vendor']) if it['vendor'] else last   # dòng thiếu hãng -> theo dòng trước
        if v:
            last = v
            if v not in order:
                order.append(v)
            cnt[v] = cnt.get(v, 0) + 1
            amt[v] = amt.get(v, 0) + (it['amount'] or 0)
    bom_total = sum(it['amount'] or 0 for it in items)
    unpriced = sum(1 for it in items if it['amount'] is None)
    pool = [v for v in order if v != 'Khác'] or order
    if pool:
        main = max(pool, key=lambda v: (amt.get(v, 0), cnt.get(v, 0)))
        how = 'theo giá trị BOM lớn nhất' if amt.get(main, 0) > 0 else 'theo hãng có nhiều dòng BOM nhất (BOM không có giá)'
        rest = sorted([v for v in order if v != main], key=lambda v: -amt.get(v, 0))
        return main, [main] + rest, how, bom_total, unpriced
    col = canon_vendor(row['Tên hãng']) if isinstance(row['Tên hãng'], str) else None
    if col and col != 'Khác':
        return col, [col], 'từ cột "Tên hãng" của CRM (BOM không ghi hãng)', bom_total, unpriced
    text = ' '.join(str(row[c]) for c in ['Tên thương vụ', 'Thông tin BOM (model, số lượng, năm support...)', 'Tên dự án']
                    if isinstance(row[c], str))
    found = sorted((m.start(), v) for v, rx in VENDOR_KEYWORDS for m in [re.search(rx, fold(text))] if m)
    if found:
        vs = [v for _, v in found]
        return vs[0], vs, 'suy ra từ từ khóa trong tên thương vụ/BOM text', bom_total, unpriced
    if col == 'Khác':
        return 'Khác', ['Khác'], 'từ cột "Tên hãng" của CRM', bom_total, unpriced
    return None, [], None, bom_total, unpriced


# ============================== DANH MỤC SALE ==============================
# Tài khoản CRM -> tên đầy đủ. Sale mới: thêm 1 dòng vào đây.
SALE_NAMES = {
    'thuydinh': 'Đinh Văn Thủy', 'hangnguyen': 'Nguyễn Thu Hằng', 'trinhluong': 'Lương Thị Tuyết Trinh',
    'phamquyet': 'Phạm Văn Quyết', 'theduong': 'Dương Văn Thế', 'quyvy': 'Vy Công Quý',
    'kienkhong': 'Khổng Đức Kiên', 'nhungnguyen': 'Nguyễn Hồng Nhung', 'hungle': 'Lê Huy Hùng',
    'anhle': 'Lê Tuấn Anh', 'thangnguyen': 'Nguyễn Duy Thắng', 'nghiemphan': 'Phan Quế Nghiệm',
    'hungpham': 'Phạm Quốc Hùng', 'tupham': 'Phạm Cao Anh Tú', 'hieudinh': 'Đinh Văn Hiệu',
    'hiendoan': 'Đoàn Thị Hiền', 'hungtran': 'Trần Quang Hưng',
}


def _sale_tags():
    """Các kiểu sale hay ghi kèm vào tên KH: _KienKhong, _QuyetPham, _Tuấn Anh, _Lê Hùng, _Trinh…"""
    tags = set()
    for user, full in SALE_NAMES.items():
        p = fold(full).split()
        fam, mid, giv = p[0], p[1:-1], p[-1]
        tags |= {user, fam + giv, giv + fam, giv, ''.join(p), ''.join(mid[-1:]) + giv}
        if len(p) >= 2:
            tags.add(''.join(p[-2:]))
    return {t for t in tags if len(t) >= 4}


SALE_TAGS = None


def strip_sale_tag(name):
    m = re.match(r'^(.*\S)\s*_\s*([^_]+)$', name)
    global SALE_TAGS
    SALE_TAGS = SALE_TAGS or _sale_tags()
    if m and re.sub(r'[^a-z]', '', fold(m.group(2))) in SALE_TAGS:
        return m.group(1).strip(), m.group(2).strip()
    return name, None


# ============================== PHÂN NHÓM KHÁCH HÀNG ==============================
SEGMENTS = ['Cơ quan Đảng, Nhà nước', 'Bộ, Ngành, Cơ quan TW', 'UBND / Sở ban ngành địa phương',
            'An ninh - Quốc phòng', 'Ngân hàng - Tài chính - Bảo hiểm', 'Viễn thông - CNTT',
            'Năng lượng - Điện - Dầu khí', 'Y tế - Giáo dục', 'Sản xuất - Thương mại - Dịch vụ',
            'Giao thông - Vận tải - Logistics', 'DNNN khác', 'Khác']
# Quy tắc áp dụng theo THỨ TỰ, khớp quy tắc đầu tiên thì dừng (so trên tên KH đã bỏ dấu, chữ thường).
SEGMENT_RULES = [
    ('Viễn thông - CNTT', r'viettel|\bvtg\b|natcom'),                       # Viettel xét trước (xem ngoại lệ bên dưới)
    ('An ninh - Quốc phòng', r'cong an|\bbca\b|\bbo ca\b|^ca |\bca (tp|tinh|thanh pho)\b|\bpa0\d|'
                             r'\b(a0[1-9]|c0[1-9]|c1[0-2]|c500|h0[1-9]|v0[1-9]|b01|x01)\b|quoc phong|\bbqp\b|quan doi|'
                             r'quan khu|bo tu lenh|binh chung|bien phong|canh sat|quan y|hai quan|military|'
                             r'tong cuc (2|ii|tinh bao|ky thuat|chinh tri)|co yeu|si quan|tong tham muu|tac chien|'
                             r'an ninh nhan dan|bo tham muu|bo doi'),
    ('Cơ quan Đảng, Nhà nước', r'tinh uy|thanh uy|quan uy|huyen uy|\bvptu\b|trung uong dang|tw dang|vptw|'
                               r'tap chi cong san|phu chu tich|van phong chu tich nuoc|quoc hoi|toa an|vien kiem sat|'
                               r'kiem toan nha nuoc'),
    ('UBND / Sở ban ngành địa phương', r'\bubnd\b|uy ban nhan dan|^so |\bso (khoa|kh|thong tin|tttt|4t|xay dung|'
                                       r'dan toc|tai chinh|y te|giao duc|noi vu|tu phap|ke hoach|cong thuong|van hoa)\b|'
                                       r'bqlda|ban quan ly khu kinh te|ban duy tu|chuyen doi so|'
                                       r'phuc vu hanh chinh cong|dai phat thanh|bao va phat thanh'),
    ('Bộ, Ngành, Cơ quan TW', r'^bo |\bbo (noi vu|tai chinh|tai nguyen|ngoai giao|dan toc|xay dung|khoa hoc|'
                              r'cong thuong|y te|giao duc|tu phap|ke hoach|nong nghiep|giao thong|van hoa|thong tin)\b|'
                              r'\bcuc\b|tong cuc|kho bac|bao hiem xa hoi|\bbhxh\b|ngan hang nha nuoc|'
                              r'(uy ban|ub) chung khoan|uy ban tieu chuan|quy phat trien|thong tin tin dung|'
                              r'\bcic\b|vnnic|ttxvn|thong tan xa|vien han lam'),
    ('Năng lượng - Điện - Dầu khí', r'\bevn|dien luc|thuy dien|nhiet dien|truyen tai dien|he thong dien|'
                                    r'thi truong dien|dieu do|\bnpt\b|\bptc ?\d|^pc |genco|\bpvn\b|pv ?gas|petro|'
                                    r'dau khi|khi viet nam|^than |\btkv\b|vimico|khoang san|nhien lieu|xang dau|nsrp|'
                                    r'npcit|\bnpc\b|\bspc\b'),
    ('Y tế - Giáo dục', r'benh vien|^bv|\bbvdk\b|y te|huyet hoc|^truong |dai hoc|hoc vien|giao duc|educa|'
                        r'\bttyt\b|hospital'),
    ('Ngân hàng - Tài chính - Bảo hiểm', r'ngan hang|bank|chung khoan|securities|bao hiem|insurance|tai chinh vi mo|'
                                         r'finance|napas|vnpay|thanh toan|\b(bidv|vietcombank|vietinbank|agribank|'
                                         r'mbbank|tpbank|vib|scb|shb|shbfc|lpb|lpbs|exim|eximbank|tcbs|hsc|kbsv|'
                                         r'vndirect|mbs|vps|shs|dnse|dsc|agriseco|opes|pjico|pvi|insmart|mfinance|'
                                         r'timo|oceanbank|publicbank)\b'),
    ('Giao thông - Vận tải - Logistics', r'hang khong|airline|airport|san bay|\bacv\b|vaeco|viags|sasco|\bvna\b|'
                                         r'truc thang|\bvnh\b|quan ly bay|vnaic|hang hai|giao hang|\bghtk\b|'
                                         r'logistics|van tai|duong sat|\bvetc\b|\bitl\b|vnpost|buu chinh'),
    ('Viễn thông - CNTT', r'vnpt|mobifone|\bcmc\b|telecom|vien thong|\bfpt\b|\bfis\b|\bctin\b|cong nghe|technology|'
                          r'tecapro|gosu|vtvcab|software|vietbay|tntech'),
    ('DNNN khác', r'cap nuoc|sawaco|vietlott|xo so|trac dia|dong tau'),
    ('Sản xuất - Thương mại - Dịch vụ', r'cong ty|\bcty\b|\bctcp\b|tnhh|tap doan|group|corporation|\bcorp\b|'
                                        r'limited|\bjsc\b|holdings'),
]
# Ngoại lệ đã xác định cụ thể (tên KH bỏ dấu, chữ thường -> nhóm). Thêm KH mới khó phân loại vào đây.
SEGMENT_OVERRIDES = {
    'tong cong ty cp buu chinh viettel': 'Giao thông - Vận tải - Logistics',
    'viettel post': 'Giao thông - Vận tải - Logistics',
    'cong ty thuong mai & xuat nhap khau/ thuong mai & xnk viettel (viettel commerce)': 'Sản xuất - Thương mại - Dịch vụ',
    'cong ty lien doanh thap ngan hang dau tu va phat trien viet nam': 'Sản xuất - Thương mại - Dịch vụ',
    'vietinbank gold & jewellery': 'Ngân hàng - Tài chính - Bảo hiểm',
    'pvi': 'Ngân hàng - Tài chính - Bảo hiểm',
    'cong ty tnhh mot thanh vien dong tau hong ha': 'DNNN khác',
    'cong ty tnhh mtv trac dia ban do': 'DNNN khác',
    'tap doan phenikaa': 'Sản xuất - Thương mại - Dịch vụ',
    'access trade': 'Sản xuất - Thương mại - Dịch vụ', 'career viet': 'Sản xuất - Thương mại - Dịch vụ',
    'diana unicharm': 'Sản xuất - Thương mại - Dịch vụ', 'everland van don': 'Sản xuất - Thương mại - Dịch vụ',
    'fujikin': 'Sản xuất - Thương mại - Dịch vụ', 'freshmart (c.p)': 'Sản xuất - Thương mại - Dịch vụ',
    'fujimart': 'Sản xuất - Thương mại - Dịch vụ', 'jarllytec': 'Sản xuất - Thương mại - Dịch vụ',
    'maison': 'Sản xuất - Thương mại - Dịch vụ', 'masterise': 'Sản xuất - Thương mại - Dịch vụ',
    'messer viet nam': 'Sản xuất - Thương mại - Dịch vụ', 'nutifood': 'Sản xuất - Thương mại - Dịch vụ',
    'panasonic r&d': 'Sản xuất - Thương mại - Dịch vụ', 'rang dong': 'Sản xuất - Thương mại - Dịch vụ',
    'shemar power vn': 'Sản xuất - Thương mại - Dịch vụ', 'sun group': 'Sản xuất - Thương mại - Dịch vụ',
    'yakult viet nam': 'Sản xuất - Thương mại - Dịch vụ', 'yokohama tyre viet nam': 'Sản xuất - Thương mại - Dịch vụ',
    'berjaya gia thinh': 'Sản xuất - Thương mại - Dịch vụ',
    'viet duc co so 3': 'Y tế - Giáo dục',
    'cong ty tnhh mtv van hanh htd va ttd quoc gia': 'Năng lượng - Điện - Dầu khí',
    'gtb endpoint protector': 'Khác', 'giai phap attt cap do 3': 'Khác',
}
# KH chưa đủ thông tin để phân nhóm chắc chắn -> để "Khác" và gắn cờ cần kiểm tra
SEGMENT_UNSURE = {'hal vn', 'mcst', 'itdb', 'lpex', 'plc', 'intech',
                  'trung tam nghien cuu va ung dung cong nghe truyen thong (r&d)'}


def classify_customer(name):
    """Trả về (nhóm KH, chắc chắn?)"""
    f = re.sub(r'\s+', ' ', fold(name)).strip()
    if f in SEGMENT_UNSURE:
        return ('Viễn thông - CNTT' if 'cong nghe' in f else 'Khác'), False
    if f in SEGMENT_OVERRIDES:
        return SEGMENT_OVERRIDES[f], True
    for seg, rx in SEGMENT_RULES:
        if re.search(rx, f):
            if seg == 'Viễn thông - CNTT' and re.search(r'viettel', f):
                if re.search(r'buu chinh|post', f):
                    return 'Giao thông - Vận tải - Logistics', True
                if re.search(r'commerce|thuong mai', f):
                    return 'Sản xuất - Thương mại - Dịch vụ', True
            return seg, True
    return 'Khác', False


# ============================== XỬ LÝ CHÍNH ==============================
def convert(src, dst, ref_date):
    raw = pd.read_excel(src, sheet_name=0)
    missing = [c for c in REQUIRED if c not in raw.columns]
    if missing:
        sys.exit(f'File raw thiếu cột: {missing}. Kiểm tra lại cấu hình export CRM.')
    if raw['ID'].duplicated().any():
        sys.exit('File raw có ID trùng – dừng để kiểm tra.')

    df = raw.copy()
    for c in df.columns:
        if df[c].dtype == object or str(df[c].dtype) in ('str', 'string'):
            df[c] = df[c].map(clean_text)
    for c in ['Timeline dự án', 'Ngày ký hợp đồng', 'Từ lúc', 'Cập nhật lần cuối']:
        parsed = df[c].map(parse_date)
        bad = df[c].notna() & parsed.isna()
        if bad.any():
            print(f'CẢNH BÁO: {bad.sum()} giá trị ngày không đọc được ở cột "{c}": {df.loc[bad, c].tolist()[:5]}')
        df[c] = parsed

    # ---- Bước 1: Tên khách hàng
    cust, cust_note, test_rows = [], [], []
    for _, r in df.iterrows():
        if isinstance(r['Tên công ty'], str) and r['Tên công ty']:
            cust.append(r['Tên công ty']); cust_note.append(None)
            if TEST_NAMES.search(fold(r['Tên thương vụ'])):
                test_rows.append(r['ID'])
            continue
        c, how = customer_from_deal_name(r['Tên thương vụ'])
        if how == 'giu' and is_solution(c):
            how = 'kiem tra'
        cust.append(c)
        if how == 'test':
            test_rows.append(r['ID']); cust_note.append(None)
        elif how == 'tach':
            cust_note.append('Tên KH lấy từ tên thương vụ (CRM trống "Tên công ty"), đã bỏ phần hãng/giải pháp/reseller')
        elif how == 'kiem tra':
            cust_note.append('Tên KH lấy nguyên tên thương vụ nhưng tên giống tên dự án/giải pháp – cần kiểm tra lại')
        else:
            cust_note.append('Tên KH lấy từ tên thương vụ (CRM trống "Tên công ty")')
    stripped = [strip_sale_tag(c) for c in cust]
    df['_cust'] = [c for c, _ in stripped]
    df['_sale_tag'] = [t for _, t in stripped]
    df['_cust_note'] = cust_note
    df = df[~df['ID'].isin(test_rows)].copy()          # loại deal test

    # gộp biến thể tên KH chỉ khác hoa/thường, dấu, khoảng trắng -> chọn bản có dấu, xuất hiện nhiều nhất
    df['_key'] = df['_cust'].map(lambda s: re.sub(r'[^a-z0-9]+', ' ', fold(s)).strip())
    def pick(g):
        vc = g.value_counts()
        # ưu tiên: có dấu tiếng Việt > không viết HOA toàn bộ > xuất hiện nhiều nhất
        return sorted(vc.index, key=lambda s: (fold(s) == s.lower(), s.isupper(), -vc[s], s))[0]
    canon = df.groupby('_key')['_cust'].agg(pick)
    df['_cust_std'] = df['_key'].map(canon)

    # ---- Bước 2: Giá trị VND chuẩn
    vinfo = df.apply(vendor_info, axis=1, result_type='expand')
    vinfo.columns = ['_main', '_all', '_vhow', '_bom', '_unpriced']
    df = pd.concat([df, vinfo], axis=1)

    name_counts = df.groupby('Tên thương vụ')['ID'].apply(list)
    out = []
    for _, r in df.iterrows():
        v0, cur, bom = r['Giá trị thương vụ'], r['Tiền tệ thương vụ'], r['_bom']
        issues, notes = [], []
        # giá trị
        if (pd.isna(v0) or v0 == 0) and bom > 0:
            val = bom; notes.append(f'Giá trị lấy từ tổng BOM ({vnd(bom)}) vì giá trị CRM = 0')
        elif pd.isna(v0) or v0 == 0:
            val = None; issues.append('Thiếu giá trị thương vụ (CRM = 0, BOM không có giá)')
        elif cur == 'USD' and v0 < USD_MISLABEL_THRESHOLD:
            val = v0 * USD_RATE; notes.append(f'Quy đổi {v0:,.0f} USD x {USD_RATE:,} (tỷ giá giả định)')
        else:
            val = v0
            if cur == 'USD':
                notes.append('CRM ghi tiền tệ USD nhưng số tiền là VND – giữ số gốc, coi là VND')
        if val and bom > 0 and abs(val - bom) > BOM_TOLERANCE * bom:
            issues.append(f'Giá trị {vnd(val)} lệch tổng BOM {vnd(bom)} ({(val - bom) / bom:+.0%})')
        if r['_unpriced'] > 0:
            issues.append(f'BOM có {r["_unpriced"]} dòng chưa có giá')
        # trạng thái / giai đoạn
        st, gd = r['Trạng thái thương vụ'], r['Giai đoạn']
        if st == 'Active' and gd == 'Thực hiện hợp đồng':
            issues.append('Trạng thái Active nhưng giai đoạn CRM là "Thực hiện hợp đồng" (có thể đã Won)')
        if st == 'Won' and gd != 'Thực hiện hợp đồng':
            issues.append(f'Trạng thái Won nhưng giai đoạn CRM là "{gd}"')
        if st == 'Won' and pd.isna(r['Ngày ký hợp đồng']):
            issues.append('Won nhưng thiếu ngày ký hợp đồng')
        if st == 'Lost' and not isinstance(r['Lý do thất bại'], str):
            issues.append('Lost nhưng thiếu lý do thất bại')
        # close date
        cd = r['Timeline dự án']
        if pd.notna(cd) and st == 'Active' and cd < ref_date:
            issues.append(f'Close date {cd:%d/%m/%Y} đã qua nhưng deal vẫn Active')
        # trùng tên
        same = [i for i in name_counts[r['Tên thương vụ']] if i != r['ID']]
        if same:
            ids = ', '.join(str(i) for i in same[:5]) + (' …' if len(same) > 5 else '')
            issues.append(f'Trùng tên thương vụ với {len(same)} deal khác (ID {ids})')
        if not isinstance(r['Chủ sở hữu thương vụ'], str):
            issues.append('Thiếu Sale (owner) trong CRM')
        # thiếu trường
        miss = [lab for lab, ok in [('Close date', pd.notna(cd)), ('Reseller', isinstance(r['Reseller'], str)),
                                    ('Phân loại HĐ', isinstance(r['Phân loại HĐ'], str))] if not ok]
        if miss:
            issues.append('Thiếu thông tin: ' + ', '.join(miss))
        # hãng
        if not isinstance(r['_main'], str):   # pandas đổi None thành NaN nên phải kiểm tra kiểu
            issues.append('Chưa xác định được hãng (BOM, cột Tên hãng và tên thương vụ đều không có)')
        else:
            notes.append(f'Hãng chính xác định {r["_vhow"]}')
        # tên KH
        if isinstance(r['_cust_note'], str):
            (issues if 'cần kiểm tra' in r['_cust_note'] else notes).append(r['_cust_note'])
        if isinstance(r['_sale_tag'], str):
            notes.append(f'Đã bỏ hậu tố tên sale "_{r["_sale_tag"]}" trong tên KH')
        if r['_cust_std'] != r['_cust']:
            notes.append(f'Tên KH chuẩn hóa từ "{r["_cust"]}"')
        # deal size
        ds = deal_size(val)
        ds0 = r['Phân loại nhóm deal']
        if ds and not isinstance(ds0, str):
            notes.append('Deal size tính từ giá trị (CRM trống)')
        elif ds and ds0 != ds:
            notes.append(f'Deal size CRM "{ds0}" không khớp giá trị – đã tính lại')
        # sale
        u = r['Chủ sở hữu thương vụ']
        if isinstance(u, str) and u.strip().lower() in SALE_NAMES:
            sale = SALE_NAMES[u.strip().lower()]
        else:
            sale = u if isinstance(u, str) else None
            if isinstance(u, str):
                issues.append(f'Sale "{u}" chưa có trong bảng tên sale – giữ tên tài khoản CRM')
        # phân nhóm khách hàng
        seg, sure = classify_customer(r['_cust_std'])
        if not sure:
            issues.append(f'Phân nhóm KH "{seg}" chưa chắc chắn – cần kiểm tra')
        # lý do failed
        reasons = [x for x in [r['Lý do thất bại'], r['Failed additional comment']] if isinstance(x, str)]
        failed = ' – '.join(reasons) if reasons else None

        vendors = r['_all'] or [None]
        n = len(vendors)
        for k, ven in enumerate(vendors):
            extra = [] if k == 0 else [f'Dòng tách hãng {k + 1}/{n} của deal ID {r["ID"]}; giá trị thương vụ chỉ ghi ở dòng hãng chính ({vendors[0]})']
            if k == 0 and n > 1:
                extra = [f'Deal có {n} hãng: {", ".join(vendors)} – đã tách thành {n} dòng, giá trị ghi ở dòng này']
            text = []
            if issues:
                text.append('[VẤN ĐỀ] ' + '; '.join(issues))
            if notes or extra:
                text.append('[ĐÃ XỬ LÝ] ' + '; '.join(extra + notes))
            assert not any(' nan' in t for t in text), f'Ghi chú lỗi ở ID {r["ID"]}: {text}'
            out.append({
                'ID': r['ID'], 'Tên thương vụ': r['Tên thương vụ'], 'Tên khách hàng': r['_cust_std'],
                'Phân nhóm khách hàng': seg, 'Sale': sale, 'Reseller': r['Reseller'], 'Giai đoạn': gd, 'Trạng thái': st,
                'Giá trị thương vụ': val if k == 0 else None, 'Hãng': ven,
                'Thông tin BOM (model, số lượng, năm support...)': r['Thông tin BOM (model, số lượng, năm support...)'],
                'Phân loại HĐ': r['Phân loại HĐ'], 'Phân loại deal size': ds,
                'Close date': cd, 'Lý do Failed': failed, 'Ngày tạo': r['Từ lúc'],
                'Cập nhật lần cuối': r['Cập nhật lần cuối'], 'Cờ chất lượng': ' | '.join(text) or None,
            })
    res = pd.DataFrame(out)
    write_excel(res, dst, ref_date, len(raw), len(test_rows))
    return res, test_rows


def write_excel(res, dst, ref_date, n_raw, n_test):
    F = 'Arial'
    wb = Workbook(); ws = wb.active; ws.title = 'Data'
    thin = Side(style='thin', color='D9D9D9'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = {'ID': 10, 'Tên thương vụ': 42, 'Tên khách hàng': 38, 'Phân nhóm khách hàng': 26, 'Sale': 20, 'Reseller': 18, 'Giai đoạn': 20,
              'Trạng thái': 11, 'Thông tin BOM (model, số lượng, năm support...)': 50,
              'Giá trị thương vụ': 18, 'Hãng': 15, 'Phân loại HĐ': 13, 'Phân loại deal size': 12, 'Close date': 12,
              'Lý do Failed': 40, 'Ngày tạo': 16, 'Cập nhật lần cuối': 16, 'Cờ chất lượng': 90}
    cols = list(res.columns)
    for j, h in enumerate(cols, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor='1F4E78'); c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[L(j)].width = widths.get(h, 15)
    split_fill = PatternFill('solid', fgColor='F2F2F2')
    for i, rec in enumerate(res.itertuples(index=False), 2):
        is_split = isinstance(rec[-1], str) and 'Dòng tách hãng' in rec[-1]
        for j, v in enumerate(rec, 1):
            if v is not None and not isinstance(v, str) and pd.isna(v):
                v = None
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            c = ws.cell(i, j, v); c.font = Font(name=F, size=10); c.border = bd
            h = cols[j - 1]
            if h == 'Giá trị thương vụ': c.number_format = '#,##0'
            elif h == 'Close date': c.number_format = 'dd/mm/yyyy'
            elif h in ('Ngày tạo', 'Cập nhật lần cuối'): c.number_format = 'dd/mm/yyyy hh:mm'
            elif h == 'ID': c.number_format = '0'
            if is_split: c.fill = split_fill
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{L(len(cols))}{len(res) + 1}'

    g = wb.create_sheet('Huong_dan')
    lines = [
        ('HƯỚNG DẪN ĐỌC FILE', None),
        ('Ngày chốt số liệu', ref_date.strftime('%d/%m/%Y')),
        ('Số deal trong file raw', n_raw),
        ('Deal test đã loại', n_test),
        ('Số deal (ID duy nhất)', res['ID'].nunique()),
        ('Số dòng trong sheet Data', len(res)),
        ('Quy tắc đếm', 'Một deal nhiều hãng được tách thành nhiều dòng cùng ID. Đếm số deal phải dùng ID duy nhất. '
                        'Giá trị thương vụ chỉ ghi ở dòng hãng chính (dòng tách tô nền xám, giá trị để trống), nên cộng '
                        'cột Giá trị thương vụ không bị đếm trùng.'),
        ('Giá trị thương vụ', 'Đơn vị VND. Để trống = chưa có dữ liệu (không phải 0). Khi CRM = 0 thì lấy từ tổng BOM.'),
        ('Hãng', 'Dòng đầu của mỗi deal là hãng chính (hãng có giá trị BOM lớn nhất). Để trống = chưa xác định được.'),
        ('Phân loại deal size', '< 1B: dưới 1 tỷ · 1B - 2B: 1 đến dưới 2 tỷ · 2B - 5B: 2 đến dưới 5 tỷ · > 5B: từ 5 tỷ. '
                                'Tính lại từ Giá trị thương vụ, không dùng phân loại CRM vì CRM ghi không nhất quán.'),
        ('Sale', 'Đổi tài khoản CRM (cột "Chủ sở hữu thương vụ") sang họ tên theo bảng ở sheet Danh_muc. '
                 'Tài khoản chưa có trong bảng thì giữ nguyên và gắn cờ.'),
        ('Phân nhóm khách hàng', 'Phân tự động theo tên KH với 12 khối ở sheet Danh_muc. Cơ quan nhà nước phân theo cấp '
                                 '(Đảng/NN, Bộ ngành TW, địa phương); mọi đơn vị thuộc Bộ Công an, Bộ Quốc phòng (kể cả bệnh viện, '
                                 'trường) vào An ninh - Quốc phòng; doanh nghiệp phân theo ngành kinh doanh chính, công ty con theo '
                                 'ngành của tập đoàn (EVN → Năng lượng, Viettel → Viễn thông, trừ Viettel Post → Logistics); '
                                 'bệnh viện/trường dân sự vào Y tế - Giáo dục. KH chưa phân được hoặc chưa chắc chắn có cờ cần kiểm tra.'),
        ('Giai đoạn', 'Giai đoạn dự án theo CRM: Đăng kí cơ hội → Tư vấn → Phương án kinh doanh → Xây dựng hồ sơ → '
                      'Thực hiện hợp đồng. Giữ nguyên giá trị CRM; mâu thuẫn với Trạng thái được ghi ở Cờ chất lượng.'),
        ('Thông tin BOM', 'Lấy nguyên cột "Thông tin BOM (model, số lượng, năm support...)" của CRM (đã bỏ ký tự HTML, '
                          'khoảng trắng thừa). Cột này do sale nhập tay, phần lớn deal để trống.'),
        ('Close date', 'Lấy từ cột "Timeline dự án" của CRM.'),
        ('Lý do Failed', 'Gộp "Lý do thất bại" và "Failed additional comment": <lý do> – <ghi chú>.'),
        ('Cờ chất lượng', '[VẤN ĐỀ] = dữ liệu thiếu/mâu thuẫn cần người phụ trách kiểm tra. '
                          '[ĐÃ XỬ LÝ] = ghi lại thay đổi so với CRM (tên KH, hãng, giá trị, deal size). Trống = không có ghi chú.'),
        ('Tham số', f'Ngưỡng USD nhập nhầm: {USD_MISLABEL_THRESHOLD:,}; tỷ giá giả định: {USD_RATE:,}; '
                    f'ngưỡng lệch BOM: {BOM_TOLERANCE:.0%}.'),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ca = g.cell(i, 1, a); ca.font = Font(name=F, size=10, bold=True)
        cb = g.cell(i, 2, b); cb.font = Font(name=F, size=10); cb.alignment = Alignment(wrap_text=True, vertical='top')
        ca.alignment = Alignment(vertical='top')
    g['A1'].font = Font(name=F, size=12, bold=True)
    g.column_dimensions['A'].width = 26; g.column_dimensions['B'].width = 110

    dm = wb.create_sheet('Danh_muc')
    hdr = PatternFill('solid', fgColor='1F4E78')
    for j, h in enumerate(['Tài khoản CRM', 'Tên Sale', '', 'Phân nhóm khách hàng', 'Số deal', 'Ví dụ khách hàng'], 1):
        if h:
            c = dm.cell(1, j, h); c.font = Font(name=F, bold=True, color='FFFFFF', size=10); c.fill = hdr
    for i, (u, n) in enumerate(SALE_NAMES.items(), 2):
        dm.cell(i, 1, u).font = Font(name=F, size=10); dm.cell(i, 2, n).font = Font(name=F, size=10)
    first = res.drop_duplicates('ID')
    for i, seg in enumerate(SEGMENTS, 2):
        sub = first[first['Phân nhóm khách hàng'] == seg]
        ex = ', '.join(sub['Tên khách hàng'].value_counts().index[:4])
        for j, v in ((4, seg), (5, len(sub)), (6, ex)):
            dm.cell(i, j, v).font = Font(name=F, size=10)
    for col, w in zip('ABCDEF', (16, 24, 3, 32, 9, 90)):
        dm.column_dimensions[col].width = w
    wb.save(dst)


def ref_from_name(path):
    m = re.search(r'_(\d{2})_(\d{2})_(\d{2})\.xlsx$', path)
    if m:
        try:
            return pd.Timestamp(datetime.strptime('/'.join(m.groups()), '%d/%m/%y'))
        except ValueError:
            pass
    return pd.Timestamp.today().normalize()


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('src'); ap.add_argument('dst'); ap.add_argument('--ref-date')
    a = ap.parse_args()
    ref = pd.Timestamp(datetime.strptime(a.ref_date, '%d/%m/%Y')) if a.ref_date else ref_from_name(a.src)
    res, tests = convert(a.src, a.dst, ref)
    print(f'Xong: {res["ID"].nunique()} deal, {len(res)} dòng, loại {len(tests)} deal test, ngày chốt {ref:%d/%m/%Y}')
